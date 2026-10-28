"""F0 tanks with no living F1, for chosen responsible people.

Automates the Ahrens Lab "F0 with No F1" list that was built by hand in
PyRAT: export each person's tanks, filter to F0, open the Pedigree
("Siblings and Children") for every F0 TID, and keep the ones with no
children. Columns match PyRAT's tank-list export so the sheet reads the same.

An F0 tank **has F1** when a crossing produced an *open* child tank and one of
the crossing's parents is

- the F0 tank itself, or
- a tank the F0's fish were split out of, with the crossing recorded no later
  than the split (followed back through earlier splits too).

Founders are usually crossed while still in their injection group tank and
only afterwards separated into their own tank (labelled "CID <crossing>-n"),
so their children hang off the group tank.

By default closed child tanks (all fish gone) do not count, so the list is
"F0s with no living F1"; ``--count-closed-f1`` counts them too. The hand-made
list was not consistent here (3 founders with only closed F1 listed, 1 not).

Checked against the hand-made October 2026 list (180 F0 tanks across 8
people): apart from that one closed-F1 case, every disagreement was an F1
born after that export or a row left out when the per-person lists were
combined.

    pixi run python -m metazebrobot.reports.f0_without_f1 ahrensm delahantyj
    pixi run python -m metazebrobot.reports.f0_without_f1 --lab ahrensm
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, BinaryIO, Callable, Dict, Iterable, List, Optional, Sequence, Set, Tuple, Union

from metazebrobot.utils.lab_time import lab_today

# PyRAT's tank-list export headers, in order.
COLUMNS = [
    "Owner", "Responsible", "Line / Strain (Name)", "Line / Strain (Official name)",
    "Rack", "Position", "Tank ID", "Generation", "Label", "Date of birth",
    "Animal age (d)", "# M", "# F", "# Total", "Age level", "Comments",
]
EVIDENCE_COLUMNS = [
    "Responsible", "Tank ID", "Line / Strain (Name)", "Crossing ID",
    "Crossing parent tank", "F1 tank IDs",
]
TANK_KEYS = [
    "tank_id", "tank_label", "tank_position", "location_rack_name",
    "owner_fullname", "responsible_id", "responsible_fullname",
    "strain_id", "strain_name", "generation", "date_of_birth", "age_level",
    "number_of_male", "number_of_female", "alive_count", "comments",
]
CROSSING_FIELDS = ["crossing_id", "date_of_record", "tanks"]
CROSSING_TANK_FIELDS = ["tank_id", "status"]
# Before the first PyRAT crossing (2019), so crossings_recorded returns all.
EARLIEST_RECORD = date(2000, 1, 1)

History = Callable[[int], List[Dict[str, Any]]]
Evidence = Tuple[int, int, List[int]]  # (crossing parent tank, crossing id, F1 tanks)


@dataclass
class F0Report:
    generated: date
    people: List[Dict[str, Any]]                   # PyRAT users, in request order
    without_f1: List[Dict[str, Any]] = field(default_factory=list)   # COLUMNS rows
    with_f1: List[Dict[str, Any]] = field(default_factory=list)      # EVIDENCE_COLUMNS rows


def crossings_by_parent(crossings: Iterable[Dict[str, Any]]) -> Dict[int, List[Dict[str, Any]]]:
    by_parent: Dict[int, List[Dict[str, Any]]] = {}
    for crossing in crossings:
        for parent in (crossing.get("tanks") or {}).get("parents") or []:
            by_parent.setdefault(parent["tank_id"], []).append(crossing)
    return by_parent


def split_sources(tank_id: int, history: Iterable[Dict[str, Any]]) -> Dict[int, str]:
    """Tanks this tank's fish were separated out of -> latest split date.

    A tank's history also carries its ancestors' events mirrored as ``copy``;
    only splits into this tank (relation None/destination) count.
    """
    sources: Dict[int, str] = {}
    for event in history:
        if event.get("event_type_name") != "SeparateEvent":
            continue
        relations = {r["tank_id"]: r.get("relation") for r in event.get("related") or []}
        if tank_id not in relations or relations[tank_id] not in (None, "destination"):
            continue
        when = (event.get("event_date") or "")[:10]
        for source, relation in relations.items():
            if relation == "source" and source != tank_id:
                sources[source] = max(sources.get(source, ""), when)
    return sources


def f1_evidence(
    tank_id: int,
    by_parent: Dict[int, List[Dict[str, Any]]],
    history: History,
    before: Optional[str] = None,
    seen: Optional[Set[int]] = None,
    count_closed: bool = False,
) -> List[Evidence]:
    """Crossings with (open, unless count_closed) child tanks this tank's fish took part in."""
    seen = seen if seen is not None else set()
    seen.add(tank_id)
    evidence: List[Evidence] = []
    for crossing in by_parent.get(tank_id, []):
        if before and (crossing.get("date_of_record") or "")[:10] > before:
            continue
        children = [
            child["tank_id"] for child in crossing["tanks"].get("children") or []
            if count_closed or child.get("status") == "open"
        ]
        if children:
            evidence.append((tank_id, crossing["crossing_id"], children))
    for source, split_date in split_sources(tank_id, history(tank_id)).items():
        if source not in seen:
            cutoff = min(split_date, before) if before else split_date
            evidence.extend(f1_evidence(source, by_parent, history, cutoff, seen, count_closed))
    return evidence


def _position_key(position: Optional[str]) -> Tuple[str, int]:
    position = position or ""
    letters = position.rstrip("0123456789")
    digits = position[len(letters):]
    return letters, int(digits) if digits else 0


def _tank_sort_key(tank: Dict[str, Any]) -> Tuple[Any, ...]:
    return (tank.get("location_rack_name") or "", _position_key(tank.get("tank_position")), tank["tank_id"])


def _birth_date(tank: Dict[str, Any]) -> Optional[date]:
    value = tank.get("date_of_birth")
    return datetime.fromisoformat(value).date() if value else None


def tank_row(tank: Dict[str, Any], official_names: Dict[int, Optional[str]], today: date) -> Dict[str, Any]:
    """One tank in PyRAT tank-list export form."""
    born = _birth_date(tank)
    comments = sorted(tank.get("comments") or [], key=lambda c: c.get("created") or "", reverse=True)
    return {
        "Owner": tank.get("owner_fullname"),
        "Responsible": tank.get("responsible_fullname"),
        "Line / Strain (Name)": tank.get("strain_name"),
        "Line / Strain (Official name)": official_names.get(tank.get("strain_id")),
        "Rack": tank.get("location_rack_name"),
        "Position": tank.get("tank_position"),
        "Tank ID": tank["tank_id"],
        "Generation": tank.get("generation"),
        "Label": tank.get("tank_label"),
        "Date of birth": born,
        "Animal age (d)": (today - born).days if born else None,
        "# M": tank.get("number_of_male"),
        "# F": tank.get("number_of_female"),
        "# Total": tank.get("alive_count"),
        "Age level": (tank.get("age_level") or "").capitalize() or None,
        "Comments": "\n\n".join(c["content"] for c in comments if c.get("content")) or None,
    }


def classify_f0_tanks(
    tanks: Iterable[Dict[str, Any]],
    by_parent: Dict[int, List[Dict[str, Any]]],
    history: History,
    official_names: Dict[int, Optional[str]],
    today: date,
    count_closed: bool = False,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Split one person's tanks into (no-F1 rows, has-F1 evidence rows)."""
    without_f1: List[Dict[str, Any]] = []
    with_f1: List[Dict[str, Any]] = []
    for tank in sorted(tanks, key=_tank_sort_key):
        if tank.get("generation") != "F0":
            continue
        evidence = f1_evidence(tank["tank_id"], by_parent, history, count_closed=count_closed)
        if not evidence:
            without_f1.append(tank_row(tank, official_names, today))
        for parent, crossing_id, children in evidence:
            with_f1.append({
                "Responsible": tank.get("responsible_fullname"),
                "Tank ID": tank["tank_id"],
                "Line / Strain (Name)": tank.get("strain_name"),
                "Crossing ID": crossing_id,
                "Crossing parent tank": parent,
                "F1 tank IDs": ", ".join(str(c) for c in children),
            })
    return without_f1, with_f1


def people_by_username(client: Any, usernames: Sequence[str]) -> List[Dict[str, Any]]:
    """PyRAT users ({userid, fullname, ...}) for these usernames, in order."""
    people = []
    for username in usernames:
        user = client.user(username)
        if user is None:
            raise ValueError(f"No PyRAT user named {username!r} (or it is not visible to this account)")
        people.append(user)
    return people


def lab_members(client: Any, owner_username: str) -> List[Dict[str, Any]]:
    """Everyone Responsible for an open tank owned by this user, by name.

    Ids come from the tanks themselves, so this works even when PyRAT's
    users endpoint does not list those people for this account.
    """
    owner = people_by_username(client, [owner_username])[0]
    tanks = client.tanks(["responsible_id", "responsible_fullname"], owner_id=owner["userid"], status=["open"])
    people = {t["responsible_id"]: t["responsible_fullname"] for t in tanks if t.get("responsible_id")}
    return [{"userid": userid, "fullname": name} for userid, name in sorted(people.items(), key=lambda p: p[1] or "")]


def build_report(
    client: Any,
    people: Sequence[Dict[str, Any]],
    today: Optional[date] = None,
    count_closed_f1: bool = False,
) -> F0Report:
    """Fetch from PyRAT and build the report for these Responsible people.

    ``people`` are PyRAT users ({userid, fullname}) from people_by_username or
    lab_members. ``client`` is a :class:`PyratApiClient` built with
    CROSSING_FIELDS and CROSSING_TANK_FIELDS.
    """
    today = today or lab_today()
    by_parent = crossings_by_parent(client.crossings_recorded(EARLIEST_RECORD))
    official_names = {s["id"]: s.get("official_name") for s in client.strains()}
    histories: Dict[int, List[Dict[str, Any]]] = {}

    def history(tank_id: int) -> List[Dict[str, Any]]:
        if tank_id not in histories:
            histories[tank_id] = client.tank_history(tank_id)
        return histories[tank_id]

    report = F0Report(generated=today, people=list(people))
    for person in people:
        tanks = client.tanks(TANK_KEYS, responsible_id=person["userid"], status=["open"])
        without_f1, with_f1 = classify_f0_tanks(
            tanks, by_parent, history, official_names, today, count_closed_f1,
        )
        report.without_f1.extend(without_f1)
        report.with_f1.extend(with_f1)
    return report


def write_xlsx(report: F0Report, target: Union[str, Path, BinaryIO]) -> None:
    """Write the report: the no-F1 list first, then the evidence for the rest."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    sheets = [
        ("Total No F1s", COLUMNS, report.without_f1),
        ("F0 with F1 (evidence)", EVIDENCE_COLUMNS, report.with_f1),
    ]
    for index, (title, columns, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([row.get(column) for column in columns])
        for (cell,) in sheet.iter_rows(min_row=2, min_col=columns.index("Tank ID") + 1, max_col=columns.index("Tank ID") + 1):
            cell.alignment = Alignment(horizontal="left")
        if "Date of birth" in columns:
            date_col = columns.index("Date of birth") + 1
            for (cell,) in sheet.iter_rows(min_row=2, min_col=date_col, max_col=date_col):
                cell.number_format = "yyyy-mm-dd"
        if "Comments" in columns:
            comment_col = columns.index("Comments") + 1
            for (cell,) in sheet.iter_rows(min_row=2, min_col=comment_col, max_col=comment_col):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            longest = max(len(str(c.value or "").split("\n")[0]) for c in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(10, longest + 2), 60)
    workbook.save(target)


def main(argv: Optional[Sequence[str]] = None) -> int:
    from metazebrobot.utils.pyrat_api_client import PyratApiClient

    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("usernames", nargs="*", help="PyRAT usernames of the Responsible people, e.g. ahrensm")
    parser.add_argument(
        "--lab", metavar="OWNER",
        help="Everyone Responsible for an open tank this username owns, e.g. ahrensm.",
    )
    parser.add_argument("-o", "--output", type=Path, help="xlsx path (default: F0 with No F1 <date>.xlsx)")
    parser.add_argument(
        "--count-closed-f1", action="store_true",
        help="Count F1 tanks that are closed (all fish gone) as F1.",
    )
    parser.add_argument("--pause", type=float, default=0.2, help="Seconds between PyRAT requests.")
    args = parser.parse_args(argv)
    if not args.usernames and not args.lab:
        parser.error("give usernames and/or --lab OWNER")

    client = PyratApiClient(CROSSING_FIELDS, CROSSING_TANK_FIELDS, pause=args.pause)
    people = people_by_username(client, args.usernames)
    if args.lab:
        named = {p["userid"] for p in people}
        people += [p for p in lab_members(client, args.lab) if p["userid"] not in named]
    report = build_report(client, people, count_closed_f1=args.count_closed_f1)
    output = args.output or Path(f"F0 with No F1 {report.generated.isoformat()}.xlsx")
    write_xlsx(report, output)

    kind = "F1" if args.count_closed_f1 else "living F1"
    for person in report.people:
        name = person["fullname"]
        missing = sum(1 for r in report.without_f1 if r["Responsible"] == name)
        having = len({r["Tank ID"] for r in report.with_f1 if r["Responsible"] == name})
        print(f"{name}: {missing} F0 tanks with no {kind}, {having} with {kind}")
    print(f"Wrote {output} ({client.requests} PyRAT requests)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
