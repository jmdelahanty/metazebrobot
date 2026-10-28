"""F0-without-F1 report rule (metazebrobot.reports.f0_without_f1).

Fixtures follow the real chain checked on 2026-10-25: founder 8515 was split
out of injection tank 6561 on 2026-03-13 and labelled "CID 17667-1"; crossing
17667 (parents 6190 x 6561, recorded 2026-03-05) raised F1 tanks 8511/8512.
8515 itself is never a crossing parent, yet the hand-made list counted it as
having F1.
"""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from metazebrobot.reports.f0_without_f1 import (
    COLUMNS,
    F0Report,
    classify_f0_tanks,
    crossings_by_parent,
    f1_evidence,
    lab_members,
    split_sources,
    tank_row,
    write_xlsx,
)


def separate(subject, source, when, others=()):
    """A SeparateEvent as served in ``subject``'s history."""
    related = [{"tank_id": source, "relation": "source"}, {"tank_id": subject, "relation": None}]
    related += [{"tank_id": t, "relation": "copy"} for t in others]
    return {"event_type_name": "SeparateEvent", "event_date": f"{when}T15:26:34", "related": related}


def crossing(crossing_id, recorded, parents, children=(), status="open"):
    return {
        "crossing_id": crossing_id,
        "date_of_record": f"{recorded}T00:00:00",
        "tanks": {
            "parents": [{"tank_id": p} for p in parents],
            "children": [{"tank_id": c, "status": status} for c in children],
        },
    }


def histories(mapping):
    return lambda tank_id: mapping.get(tank_id, [])


def test_split_sources_ignores_ancestor_events_mirrored_as_copy():
    history = [
        separate(8515, 6561, "2026-03-13"),
        separate(8515, 6561, "2026-03-14"),
        # 6561 -> 8443 split, mirrored into 8515's history.
        {"event_type_name": "SeparateEvent", "event_date": "2026-02-27T16:06:55",
         "related": [{"tank_id": 6561, "relation": "source"}, {"tank_id": 8443, "relation": None},
                     {"tank_id": 8515, "relation": "copy"}]},
    ]
    assert split_sources(8515, history) == {6561: "2026-03-14"}


def test_direct_crossing_with_open_child_is_f1():
    by_parent = crossings_by_parent([crossing(19288, "2026-10-08", [8544, 9268], [9655])])
    assert f1_evidence(9268, by_parent, histories({})) == [(9268, 19288, [9655])]


def test_closed_children_and_childless_crossings_do_not_count():
    by_parent = crossings_by_parent([
        crossing(18088, "2026-05-02", [8173, 7867], [8882], status="closed"),
        crossing(19372, "2026-10-01", [7867]),
    ])
    assert f1_evidence(7867, by_parent, histories({})) == []
    assert f1_evidence(7867, by_parent, histories({}), count_closed=True) == [(7867, 18088, [8882])]


def test_founder_inherits_crossings_of_its_group_tank_up_to_the_split():
    by_parent = crossings_by_parent([
        crossing(17667, "2026-03-05", [6190, 6561], [8511, 8512]),
        crossing(18500, "2026-05-01", [6561], [9000]),   # after 8515 left 6561
    ])
    history = histories({8515: [separate(8515, 6561, "2026-03-13")]})
    assert f1_evidence(8515, by_parent, history) == [(6561, 17667, [8511, 8512])]


def test_split_chains_are_followed_and_cycles_stop():
    by_parent = crossings_by_parent([crossing(17379, "2026-01-10", [6190, 6561], [8278])])
    history = histories({
        8285: [separate(8285, 6639, "2026-02-01")],
        6639: [separate(6639, 6561, "2025-03-05"), separate(6639, 8285, "2026-03-01")],
    })
    # 17379 was recorded after 6639 left 6561, so it is not 8285's.
    assert f1_evidence(8285, by_parent, history) == []
    history_later = histories({
        8285: [separate(8285, 6639, "2026-02-01")],
        6639: [separate(6639, 6561, "2026-01-20")],
    })
    assert f1_evidence(8285, by_parent, history_later) == [(6561, 17379, [8278])]


def tank(tank_id, generation="F0", rack="M17", position="B4", **extra):
    return {
        "tank_id": tank_id, "generation": generation, "location_rack_name": rack,
        "tank_position": position, "owner_fullname": "Ahrens Misha",
        "responsible_fullname": "Delahanty Jeremy", "strain_id": 1034,
        "strain_name": "Tg(UAS:COX4-COX8-dL5**-mCer3)", "date_of_birth": "2026-02-03T00:00:00",
        "age_level": "adult", "number_of_male": 0, "number_of_female": 0, "alive_count": 22,
        "tank_label": None, "comments": [], **extra,
    }


def test_tank_row_matches_pyrat_export_format():
    row = tank_row(tank(6073, comments=[
        {"created": "2024-10-13T10:25:00", "content": "inj w/ Tol2-gfap-Cas9-mScarlet3"},
        {"created": "2025-06-03T16:13:00", "content": "AQ-050-2"},
    ]), {1034: "Tg(official)"}, date(2026, 10, 25))
    assert list(row) == COLUMNS
    assert row["Comments"] == "AQ-050-2\n\ninj w/ Tol2-gfap-Cas9-mScarlet3"
    assert row["Line / Strain (Official name)"] == "Tg(official)"
    assert row["Date of birth"] == date(2026, 2, 3)
    assert row["Animal age (d)"] == 264
    assert row["Age level"] == "Adult"
    assert row["# Total"] == 22


def test_classify_keeps_f0_only_sorted_by_rack_and_position():
    by_parent = crossings_by_parent([crossing(1, "2026-01-01", [3], [4])])
    tanks = [
        tank(1, position="B10"), tank(2, position="B9"),
        tank(3), tank(5, generation="F1"),
    ]
    without_f1, with_f1 = classify_f0_tanks(tanks, by_parent, histories({}), {}, date(2026, 10, 25))
    assert [r["Tank ID"] for r in without_f1] == [2, 1]
    assert [(r["Tank ID"], r["Crossing ID"], r["F1 tank IDs"]) for r in with_f1] == [(3, 1, "4")]


def test_write_xlsx_round_trips_headers_and_rows():
    report = F0Report(generated=date(2026, 10, 25), people=[])
    report.without_f1 = [tank_row(tank(8349), {}, date(2026, 10, 25))]
    buffer = BytesIO()
    write_xlsx(report, buffer)
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    sheet = workbook["Total No F1s"]
    assert [c.value for c in sheet[1]] == COLUMNS
    assert sheet.cell(row=2, column=COLUMNS.index("Tank ID") + 1).value == 8349
    assert workbook.sheetnames == ["Total No F1s", "F0 with F1 (evidence)"]


class FakeClient:
    def __init__(self, users, tanks):
        self._users, self._tanks = users, tanks

    def user(self, username):
        return self._users.get(username)

    def tanks(self, keys, owner_id=None, **filters):
        return [t for t in self._tanks if t["owner_id"] == owner_id]


def test_lab_members_come_from_owned_tanks_not_the_users_endpoint():
    client = FakeClient(
        {"ahrensm": {"userid": 7, "fullname": "Ahrens Misha"}},
        [
            {"owner_id": 7, "responsible_id": 97, "responsible_fullname": "Ruetten Virginie"},
            {"owner_id": 7, "responsible_id": 76, "responsible_fullname": "How Javier"},
            {"owner_id": 7, "responsible_id": 76, "responsible_fullname": "How Javier"},
            {"owner_id": 3, "responsible_id": 5, "responsible_fullname": "Someone Else"},
        ],
    )
    assert lab_members(client, "ahrensm") == [
        {"userid": 76, "fullname": "How Javier"},
        {"userid": 97, "fullname": "Ruetten Virginie"},
    ]
